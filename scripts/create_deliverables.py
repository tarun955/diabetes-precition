"""Generate the rubric notebook and the exactly-four-slide presentation."""

from __future__ import annotations

import json
from pathlib import Path

import nbformat as nbf
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]


def build_notebook() -> None:
    output_dir = ROOT / "notebooks"
    output_dir.mkdir(exist_ok=True)
    md = lambda text: nbf.v4.new_markdown_cell(text.strip())
    code = lambda text: nbf.v4.new_code_cell(text.strip())
    cells = [
        md("""
# Diabetes Prediction Using Ensemble Learning

**Comparison:** Random Forest vs Gradient Boosting  
**Dataset:** Kaggle `iammustafatz/diabetes-prediction-dataset`

## Project objective
Predict the binary `diabetes` outcome and compare two ensemble classifiers. This notebook covers the faculty rubric: dataset understanding, preprocessing, model application, randomized tuning, evaluation, visualization, manual metric calculation, comparison, and interpretation.

> This is an educational model study, not a clinical diagnosis.
"""),
        md("""
## 1. Imports and configuration

The data file is a variable so the exact faculty CSV can be used without changing the methodology. The reusable training functions keep this notebook and the deployed API consistent.
"""),
        code("""
from pathlib import Path
import sys
import matplotlib.pyplot as plt
import pandas as pd
from sklearn.ensemble import GradientBoostingClassifier, RandomForestClassifier
from sklearn.metrics import confusion_matrix, ConfusionMatrixDisplay
from sklearn.model_selection import RandomizedSearchCV, StratifiedKFold, train_test_split

ROOT = Path.cwd()
if ROOT.name == 'notebooks':
    ROOT = ROOT.parent
sys.path.insert(0, str(ROOT))
from scripts.train_models import (
    FEATURES, TARGET, build_pipeline, clean_data, evaluate, load_data,
    resolve_data_file, save_visualizations,
)

DATA_FILE = ROOT / 'data' / 'diabetes_prediction_dataset.csv'
if not DATA_FILE.exists():
    DATA_FILE = resolve_data_file(None)
if DATA_FILE is None:
    raise FileNotFoundError('Add the faculty CSV to data/ or update DATA_FILE.')
print(f'Using: {DATA_FILE}')
"""),
        md("""
## 2. Dataset understanding

The dataset contains demographic, lifestyle, and clinical measurements. The target is `diabetes`: 0 means no recorded diabetes and 1 indicates diabetes. The next cell reports record count, data types, missing values, target balance, and a sample from the supplied file.
"""),
        code("""
raw_df, source = load_data(DATA_FILE, '')
df = clean_data(raw_df)
print('Source:', source)
print('Shape:', df.shape)
display(df.head())
display(pd.DataFrame({'dtype': df.dtypes.astype(str), 'missing': df.isna().sum(), 'unique': df.nunique()}))
display(df[TARGET].value_counts().rename('records').to_frame())
"""),
        md("""
## 3. Preprocessing decisions

1. Numeric values are converted to numeric types and missing values use the training median.
2. Categorical values use the most frequent category and one-hot encoding.
3. Numeric fields are standardized inside the pipeline for a reproducible data contract.
4. IQR clipping reduces extreme data-entry influence; limits are learned only on training data.
5. Stratification keeps both target classes represented in the split and cross-validation.
"""),
        code("""
numeric_columns = ['age', 'hypertension', 'heart_disease', 'bmi', 'HbA1c_level', 'blood_glucose_level']
categorical_columns = ['gender', 'smoking_history']
print('Numeric features:', numeric_columns)
print('Categorical features:', categorical_columns)
display(df[FEATURES].isna().sum().to_frame('missing before pipeline imputation'))
"""),
        md("""
## 4. Model application and parameter tuning

Random Forest uses many independent trees and averages their votes. Gradient Boosting grows trees sequentially; each new tree focuses on the current errors. `RandomizedSearchCV` samples model-specific hyperparameters and scores each configuration with ROC-AUC.
"""),
        code("""
X = df[FEATURES]
y = df[TARGET]
X_train, X_test, y_train, y_test = train_test_split(X, y, test_size=0.20, stratify=y, random_state=42)
cv = StratifiedKFold(n_splits=3, shuffle=True, random_state=42)

search_spaces = {
    'random_forest': (
        build_pipeline(RandomForestClassifier(random_state=42, n_jobs=-1)),
        {'model__n_estimators': [150, 250, 350], 'model__max_depth': [None, 8, 12, 18],
         'model__min_samples_leaf': [1, 2, 4, 8], 'model__max_features': ['sqrt', 'log2', 0.7],
         'model__class_weight': [None, 'balanced', 'balanced_subsample']},
    ),
    'gradient_boosting': (
        build_pipeline(GradientBoostingClassifier(random_state=42)),
        {'model__n_estimators': [100, 150, 250], 'model__learning_rate': [0.03, 0.05, 0.1, 0.15],
         'model__max_depth': [2, 3, 4], 'model__min_samples_leaf': [1, 3, 6], 'model__subsample': [0.8, 1.0]},
    ),
}

searches = {}
evaluations = {}
for name, (pipeline, params) in search_spaces.items():
    search = RandomizedSearchCV(pipeline, params, n_iter=8, scoring='roc_auc', cv=cv, n_jobs=-1, random_state=42, refit=True)
    search.fit(X_train, y_train)
    searches[name] = search
    evaluations[name] = evaluate(name, search, X_test, y_test)
    print(name, evaluations[name]['best_params'])
"""),
        md("""
## 5. Evaluation metrics

Accuracy measures total correctness. Precision asks how many predicted positives are truly positive. Recall asks how many actual positives were detected. F1 balances precision and recall. ROC-AUC summarizes probability ranking quality. The confusion matrix makes false positives and false negatives visible.
"""),
        code("""
metrics_table = pd.DataFrame({name: result['metrics'] for name, result in evaluations.items()}).T
display(metrics_table.style.format('{:.3f}'))
for name, result in evaluations.items():
    print(f'\\n{name.replace("_", " ").title()} manual metrics:')
    print(result['manual_metrics'])
    ConfusionMatrixDisplay.from_predictions(y_test, result['predictions'], cmap='Blues')
    plt.title(name.replace('_', ' ').title())
    plt.show()
"""),
        md("""
### Manual calculation example

For a binary confusion matrix with TN, FP, FN, and TP:

- `accuracy = (TP + TN) / (TP + TN + FP + FN)`
- `precision = TP / (TP + FP)`
- `recall = TP / (TP + FN)`

The next cell performs this calculation independently of scikit-learn's convenience functions.
"""),
        code("""
rf_predictions = evaluations['random_forest']['predictions']
tn, fp, fn, tp = confusion_matrix(y_test, rf_predictions, labels=[0, 1]).ravel()
total = tn + fp + fn + tp
manual_accuracy = (tp + tn) / total
manual_precision = tp / (tp + fp) if (tp + fp) else 0
manual_recall = tp / (tp + fn) if (tp + fn) else 0
print({'TN': tn, 'FP': fp, 'FN': fn, 'TP': tp})
print({'accuracy': manual_accuracy, 'precision': manual_precision, 'recall': manual_recall})
"""),
        md("""
## 6. Customized visualizations and feature interpretation

The plots are created from this dataset and the held-out predictions: target balance, target correlations, confusion matrices, ROC curves, and metric comparison. Feature importance shows which transformed signals each ensemble used most heavily. It indicates model reliance, not causation.
"""),
        code("""
save_visualizations(df, y_test, evaluations)
for name, result in evaluations.items():
    print(name, result['feature_importance'][:8])
print('Saved customized plots to outputs/')
"""),
        md("""
## 7. Comparison and conclusion

Compare the two rows in the metric table rather than relying on accuracy alone. If Gradient Boosting leads on ROC-AUC/F1, sequential error-correction may be capturing useful interactions. If Random Forest is close or better, averaging may be generalizing more stably for this dataset.

For the final report, state the actual winning model, quote the actual values, discuss the precision–recall trade-off, and mention that a benchmark probability is not a diagnosis.

### Rubric checklist
- Dataset features, target, records, missing values, and data types: sections 1–2
- Preprocessing and justification: section 3
- Algorithm steps and application: section 4
- Randomized parameter tuning: section 4
- Metrics and manual calculation: section 5
- Comparison, visuals, interpretation, and conclusion: sections 5–7
"""),
    ]
    notebook = nbf.v4.new_notebook(cells=cells, metadata={
        "kernelspec": {"display_name": "Python 3", "language": "python", "name": "python3"},
        "language_info": {"name": "python", "version": "3.12"},
    })
    nbf.write(notebook, output_dir / "diabetes_ensemble_study.ipynb")


COLORS = {
    "navy": RGBColor(23, 45, 58), "teal": RGBColor(25, 139, 131),
    "mint": RGBColor(231, 246, 243), "coral": RGBColor(239, 124, 99),
    "ink": RGBColor(44, 61, 72), "muted": RGBColor(111, 126, 136),
    "line": RGBColor(220, 229, 231), "white": RGBColor(255, 255, 255),
}


def text(slide, value, x, y, w, h, size=16, color="ink", bold=False, align=PP_ALIGN.LEFT):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = shape.text_frame
    frame.clear(); frame.margin_left = Inches(.02); frame.margin_right = Inches(.02); frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = frame.paragraphs[0]; paragraph.alignment = align
    run = paragraph.add_run(); run.text = value; run.font.name = "Aptos"; run.font.size = Pt(size); run.font.bold = bold; run.font.color.rgb = COLORS[color]
    return shape


def box(slide, x, y, w, h, fill="white", line="line"):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid(); shape.fill.fore_color.rgb = COLORS[fill]; shape.line.color.rgb = COLORS[line]; shape.line.width = Pt(.7)
    return shape


def footer(slide, number, color="muted"):
    text(slide, "DIABETES ENSEMBLE LAB  •  MINI-PROJECT", .55, 7.08, 5.5, .25, 7, color, True)
    text(slide, f"0{number} / 04", 11.9, 7.08, .85, .25, 7, color, True, PP_ALIGN.RIGHT)


def metric_value(metrics, model, key):
    value = metrics.get(model, {}).get(key)
    return "—" if value is None else f"{value:.3f}"


def build_presentation() -> None:
    output_dir = ROOT / "presentation"
    output_dir.mkdir(exist_ok=True)
    metadata_file = ROOT / "models" / "metadata.json"
    metadata = json.loads(metadata_file.read_text(encoding="utf-8")) if metadata_file.exists() else {}
    metrics = metadata.get("metrics", {})
    prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)

    slide = prs.slides.add_slide(prs.slide_layouts[6]); slide.background.fill.solid(); slide.background.fill.fore_color.rgb = COLORS["navy"]
    text(slide, "ENSEMBLE MODELING STUDY", .7, .72, 4.5, .3, 10, "mint", True)
    text(slide, "Predicting diabetes\nwith two powerful lenses.", .7, 1.45, 7.7, 1.7, 31, "white", True)
    text(slide, "Random Forest  ×  Gradient Boosting", .72, 3.5, 5.2, .4, 15, "mint")
    text(slide, "A concise comparison of preprocessing, tuning, evaluation, and real-world interpretation using the faculty-approved Kaggle dataset.", .72, 4.2, 6.0, .8, 12, "white")
    circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(8.5), Inches(1.4), Inches(3.2), Inches(3.2)); circle.fill.solid(); circle.fill.fore_color.rgb = COLORS["teal"]; circle.line.color.rgb = COLORS["teal"]
    text(slide, "RF  ×  GB", 8.9, 2.2, 2.4, .6, 25, "white", True, PP_ALIGN.CENTER); text(slide, "clinical signals\n→ model evidence", 8.95, 3.02, 2.3, .7, 11, "mint", True, PP_ALIGN.CENTER); footer(slide, 1, "muted")

    slide = prs.slides.add_slide(prs.slide_layouts[6]); slide.background.fill.solid(); slide.background.fill.fore_color.rgb = COLORS["white"]
    text(slide, "01  /  DATA + METHOD", .65, .52, 4.5, .25, 9, "teal", True); text(slide, "From patient signals to tuned ensembles.", .65, .92, 9.4, .55, 24, "navy", True)
    cards = [("08", "input features", "Age, BMI, HbA1c, glucose,\nconditions, and categories."), ("01", "binary target", "diabetes = 0 or 1"), ("CV", "randomized tuning", "Stratified search keeps\nclass balance visible.")]
    for i, (number, label, body) in enumerate(cards):
        x = .7 + i * 4.12; box(slide, x, 1.85, 3.65, 1.28, "mint", "mint"); text(slide, number, x+.2, 2.04, .65, .38, 21, "teal", True); text(slide, label.upper(), x+1, 2.03, 2.3, .25, 8, "teal", True); text(slide, body, x+1, 2.3, 2.4, .5, 10)
    text(slide, "PREPROCESSING PIPELINE", .72, 3.72, 3.2, .25, 9, "muted", True)
    steps = [("01", "Validate", "types + schema"), ("02", "Impute", "median / mode"), ("03", "Encode", "one-hot categories"), ("04", "Clip + scale", "IQR + standardize"), ("05", "Fit", "two ensembles")]
    for i, (number, title, body) in enumerate(steps):
        x = .72 + i * 2.42; text(slide, number, x, 4.17, .35, .24, 9, "teal", True); text(slide, title, x, 4.52, 1.75, .28, 12, "navy", True); text(slide, body, x, 4.88, 1.8, .35, 9, "muted")
        if i < 4: text(slide, "→", x+1.83, 4.5, .3, .3, 16, "coral", True)
    box(slide, .72, 5.62, 11.85, .62, "navy", "navy"); text(slide, "Random Forest averages independent tree votes. Gradient Boosting corrects earlier tree errors sequentially.", .95, 5.77, 11.3, .26, 11, "white", True); footer(slide, 2)

    slide = prs.slides.add_slide(prs.slide_layouts[6]); slide.background.fill.solid(); slide.background.fill.fore_color.rgb = COLORS["white"]
    text(slide, "02  /  RESULTS", .65, .52, 4.5, .25, 9, "teal", True); text(slide, "Compare more than accuracy.", .65, .92, 8, .55, 24, "navy", True); text(slide, "Held-out validation metrics after randomized hyperparameter search", .67, 1.52, 7, .3, 10, "muted")
    for x, header in zip([.85, 4.45, 8.35], ["METRIC", "RANDOM FOREST", "GRADIENT BOOSTING"]): text(slide, header, x, 2.13, 3.4, .25, 8, "muted", True)
    for row, (label, key) in enumerate([("Accuracy", "accuracy"), ("Precision", "precision"), ("Recall", "recall"), ("F1 score", "f1"), ("ROC-AUC", "roc_auc")]):
        y = 2.55 + row * .57; text(slide, label, .85, y, 2.4, .3, 11, "ink", row == 4); box(slide, 4.32, y-.04, 2.4, .38, "mint" if metrics else "white"); box(slide, 8.22, y-.04, 2.4, .38, "mint" if metrics else "white"); text(slide, metric_value(metrics, "random_forest", key), 4.32, y-.02, 2.4, .3, 13, "teal" if metrics else "muted", True, PP_ALIGN.CENTER); text(slide, metric_value(metrics, "gradient_boosting", key), 8.22, y-.02, 2.4, .3, 13, "teal" if metrics else "muted", True, PP_ALIGN.CENTER)
    box(slide, .85, 5.65, 10.95, .66, "mint", "mint"); text(slide, "Scores are populated from models/metadata.json after training." if not metrics else "Use ROC-AUC and F1 alongside recall to justify the model selected for the final demo.", 1.08, 5.84, 10.4, .25, 10, "ink", True); footer(slide, 3)

    slide = prs.slides.add_slide(prs.slide_layouts[6]); slide.background.fill.solid(); slide.background.fill.fore_color.rgb = COLORS["navy"]
    text(slide, "03  /  INTERPRETATION", .65, .52, 4.5, .25, 9, "mint", True); text(slide, "The best model is the one\nthat fits the decision.", .65, 1, 7.2, 1.25, 25, "white", True)
    blocks = [("01", "Read the trade-off", "Recall protects against missed positive cases; precision limits unnecessary alarms."), ("02", "Explain the signals", "Feature importance shows model reliance, not causation. Discuss glucose, HbA1c, BMI, and context."), ("03", "Keep the boundary clear", "A benchmark probability supports analysis. It is not a clinical diagnosis.")]
    for i, (number, title, body) in enumerate(blocks):
        y = 2.75 + i * 1; text(slide, number, .75, y, .4, .3, 10, "coral", True); text(slide, title, 1.45, y, 3, .3, 12, "white", True); text(slide, body, 4.1, y, 7.8, .5, 10, "mint")
    box(slide, .7, 6, 11.7, .56, "teal", "teal"); text(slide, "Notebook  •  tuned artifacts  •  Vercel UI  •  reproducible comparison", .95, 6.16, 11.2, .24, 10, "white", True, PP_ALIGN.CENTER); footer(slide, 4, "muted")
    prs.save(output_dir / "diabetes_ensemble_presentation.pptx")


if __name__ == "__main__":
    build_notebook()
    # The final deck uses the real working prototype screenshots. Keeping the
    # deck builder separate prevents the notebook generation code from making
    # assumptions about presentation layout.
    from create_screenshot_deck import build as build_screenshot_deck

    build_screenshot_deck()
    print("Created notebook and four-slide screenshot presentation.")
