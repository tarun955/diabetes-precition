"""Create the four-slide prototype presentation with real app screenshots."""

from __future__ import annotations

from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
SCREENSHOTS = ROOT / "presentation" / "screenshots"
OUTPUT = ROOT / "presentation" / "diabetes_ensemble_presentation.pptx"

COLORS = {
    "paper": RGBColor(229, 232, 225),
    "ink": RGBColor(23, 47, 45),
    "muted": RGBColor(53, 81, 76),
    "teal": RGBColor(5, 104, 95),
    "dark": RGBColor(18, 63, 59),
    "coral": RGBColor(168, 68, 63),
    "ochre": RGBColor(229, 204, 112),
    "white": RGBColor(247, 250, 245),
    "line": RGBColor(184, 197, 188),
}


def add_text(slide, value: str, x: float, y: float, w: float, h: float, *, size: int, color: str = "ink", bold: bool = False, font: str = "Aptos", align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.margin_left = Inches(0.02)
    frame.margin_right = Inches(0.02)
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = value
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = COLORS[color]
    return box


def add_block(slide, x: float, y: float, w: float, h: float, fill: str, line: str | None = None):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = COLORS[fill]
    shape.line.color.rgb = COLORS[line or fill]
    shape.line.width = Pt(0.6)
    return shape


def add_header(slide, number: int, topic: str):
    add_text(slide, "DIABETES MODEL STUDY", 0.55, 0.28, 3.2, 0.2, size=8, color="teal", bold=True)
    add_text(slide, topic, 9.0, 0.28, 3.5, 0.2, size=8, color="muted", align=PP_ALIGN.RIGHT)
    add_text(slide, f"{number:02d} / 04", 12.0, 7.13, 0.75, 0.18, size=8, color="muted", bold=True, align=PP_ALIGN.RIGHT)
    add_text(slide, "Working Next.js prototype · educational use only", 0.55, 7.13, 4.5, 0.18, size=8, color="muted")


def add_screenshot(slide, filename: str, x: float = 5.0, y: float = 0.9, w: float = 7.75, h: float = 5.75):
    path = SCREENSHOTS / filename
    with Image.open(path) as image:
        width, height = image.size
    ratio = width / height
    max_w, max_h = w, h
    if max_w / max_h > ratio:
        draw_h = max_h
        draw_w = draw_h * ratio
    else:
        draw_w = max_w
        draw_h = draw_w / ratio
    left = x + (max_w - draw_w) / 2
    top = y + (max_h - draw_h) / 2
    add_block(slide, left - 0.05, top - 0.05, draw_w + 0.1, draw_h + 0.1, "white", "line")
    slide.shapes.add_picture(str(path), Inches(left), Inches(top), width=Inches(draw_w), height=Inches(draw_h))


def add_bullet(slide, text: str, x: float, y: float, w: float, *, color: str = "ink", mark: str = "•"):
    add_text(slide, mark, x, y, 0.18, 0.24, size=16, color="coral", bold=True)
    add_text(slide, text, x + 0.28, y - 0.01, w - 0.28, 0.5, size=13, color=color)


def build() -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # Slide 1: overview.
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = COLORS["paper"]
    add_header(slide, 1, "Overview")
    add_text(slide, "The app starts\nwith one question.", 0.55, 1.05, 4.1, 1.25, size=29, color="ink", bold=False, font="Georgia")
    add_text(slide, "What do these models see?", 0.58, 2.54, 3.8, 0.42, size=15, color="teal", bold=True)
    add_text(slide, "The overview turns an ensemble-learning study into a visual reading: one profile, eight signals, and two different ways to combine small decisions.", 0.58, 3.05, 3.9, 1.0, size=15, color="muted", font="Georgia")
    add_block(slide, 0.58, 4.42, 3.65, 1.0, "dark")
    add_text(slide, "Both learners read the same record.\nTheir learning habits are different.", 0.82, 4.64, 3.15, 0.5, size=15, color="white", font="Georgia")
    add_screenshot(slide, "01-overview.png")

    # Slide 2: prediction flow.
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = COLORS["paper"]
    add_header(slide, 2, "Prediction")
    add_text(slide, "One profile,\ntwo estimates.", 0.55, 1.05, 4.1, 1.2, size=30, color="ink", font="Georgia")
    add_text(slide, "The working flow", 0.58, 2.62, 3.7, 0.32, size=16, color="teal", bold=True)
    add_bullet(slide, "Enter eight fields from the dataset.", 0.58, 3.10, 3.85)
    add_bullet(slide, "Send the same profile to both trained models.", 0.58, 3.82, 3.85)
    add_bullet(slide, "Read lower or higher likelihood in plain language.", 0.58, 4.54, 3.85)
    add_block(slide, 0.58, 5.52, 3.65, 0.72, "ochre")
    add_text(slide, "The percentage is a model output — not a diagnosis or a personal medical risk score.", 0.8, 5.69, 3.18, 0.36, size=12, color="ink", font="Georgia")
    add_screenshot(slide, "02-predict-result.png")

    # Slide 3: comparison.
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = COLORS["paper"]
    add_header(slide, 3, "Compare models")
    add_text(slide, "A score answers\na question.", 0.55, 1.05, 4.1, 1.2, size=30, color="ink", font="Georgia")
    add_text(slide, "Why the table matters", 0.58, 2.62, 3.8, 0.32, size=16, color="teal", bold=True)
    add_text(slide, "Accuracy asks how often the whole prediction was correct. Precision asks how often a positive flag was right. Recall asks how many recorded positive cases were found.", 0.58, 3.12, 3.95, 1.2, size=14, color="muted", font="Georgia")
    add_block(slide, 0.58, 4.75, 3.65, 1.08, "sky" if False else "white")
    add_text(slide, "In this held-out test, Gradient Boosting edges ahead on ROC-AUC and recall. Random Forest has the higher precision.", 0.82, 4.96, 3.12, 0.65, size=13, color="ink", font="Georgia")
    add_screenshot(slide, "03-compare-models.png")

    # Slide 4: data and boundaries.
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = COLORS["paper"]
    add_header(slide, 4, "Understand the data")
    add_text(slide, "What the models\nare allowed to see.", 0.55, 1.05, 4.1, 1.2, size=29, color="ink", font="Georgia")
    add_text(slide, "The study context", 0.58, 2.62, 3.8, 0.32, size=16, color="teal", bold=True)
    add_bullet(slide, "100,000 records and 20,000 held out for testing.", 0.58, 3.10, 3.85)
    add_bullet(slide, "Eight inputs: history, categories, BMI, HbA1c, and glucose.", 0.58, 3.82, 3.85)
    add_bullet(slide, "The pipeline handles categories, missing values, scale, and extremes.", 0.58, 4.54, 3.85)
    add_block(slide, 0.58, 5.52, 3.65, 0.72, "dark")
    add_text(slide, "Feature importance describes model reliance. It does not prove cause.", 0.8, 5.69, 3.18, 0.36, size=12, color="white", font="Georgia")
    add_screenshot(slide, "04-understand-the-data.png")

    prs.save(OUTPUT)
    print(f"Created {OUTPUT} with {len(prs.slides)} slides")


if __name__ == "__main__":
    build()
